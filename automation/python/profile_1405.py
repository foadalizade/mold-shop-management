import json
import os
from collections import Counter
from datetime import date, datetime, time

import openpyxl
from openpyxl.utils import get_column_letter

SOURCE = "data/source/1405.xlsx"
OUT_JSON = "data/processed/1405_profile.json"
OUT_MD = "data/processed/1405_profile.md"


def typ(v):
    if v in (None, ""):
        return "blank"
    if isinstance(v, bool):
        return "boolean"
    if isinstance(v, datetime):
        return "datetime"
    if isinstance(v, date):
        return "date"
    if isinstance(v, time):
        return "time"
    if isinstance(v, (int, float)):
        return "number"
    return "text"


def clean(v):
    return v.isoformat() if isinstance(v, (datetime, date, time)) else v


# IMPORTANT:
# read_only=True streams the workbook instead of materializing the entire
# worksheet in memory. This is much faster and safer for the large 1405.xlsx.
wb = openpyxl.load_workbook(
    SOURCE,
    read_only=True,
    data_only=False,
)

out = {
    "generated_at_utc": datetime.utcnow().isoformat() + "Z",
    "source_file": SOURCE,
    "sheet_count": len(wb.sheetnames),
    "sheets": [],
}

for ws in wb.worksheets:
    mr = ws.max_row or 0
    mc = ws.max_column or 0

    # Per-column statistics are collected in ONE streaming pass.
    nonblank = [0] * mc
    type_counts = [Counter() for _ in range(mc)]
    samples = [[] for _ in range(mc)]
    seen = [set() for _ in range(mc)]

    headers = []
    header_row = None
    nonempty_rows = 0
    formula_count = 0
    first_6_rows = []

    for row_idx, row in enumerate(ws.iter_rows(), start=1):
        row_has_value = False
        row_values = []

        for col_idx, cell in enumerate(row, start=1):
            v = cell.value
            row_values.append(clean(v))

            if v not in (None, ""):
                row_has_value = True
                nonblank[col_idx - 1] += 1
                type_counts[col_idx - 1][typ(v)] += 1

                if len(samples[col_idx - 1]) < 20:
                    key = str(clean(v))
                    if key not in seen[col_idx - 1]:
                        seen[col_idx - 1].add(key)
                        samples[col_idx - 1].append(key)

            if cell.data_type == "f":
                formula_count += 1

        if row_has_value:
            nonempty_rows += 1

            if header_row is None and row_idx <= 30:
                header_row = row_idx
                headers = row_values

        if len(first_6_rows) < 6:
            first_6_rows.append({
                "row": row_idx,
                "values": row_values,
            })

    merged_ranges = []
    try:
        merged_ranges = [str(x) for x in list(ws.merged_cells.ranges)[:100]]
        merged_count = len(ws.merged_cells.ranges)
    except Exception:
        merged_count = 0

    columns = []
    for idx in range(mc):
        columns.append({
            "column_index": idx + 1,
            "column_letter": get_column_letter(idx + 1),
            "provisional_header": headers[idx] if idx < len(headers) else None,
            "nonblank_count": nonblank[idx],
            "blank_count": max(0, mr - nonblank[idx]),
            "data_types": dict(type_counts[idx]),
            "sample_unique_values": samples[idx],
        })

    out["sheets"].append({
        "sheet_name": ws.title,
        "max_row": mr,
        "max_column": mc,
        "nonempty_rows": nonempty_rows,
        "formula_count": formula_count,
        "merged_range_count": merged_count,
        "merged_ranges": merged_ranges,
        "provisional_header_row": header_row,
        "provisional_headers": headers,
        "sample_rows_first_6": first_6_rows,
        "columns": columns,
    })

wb.close()

os.makedirs("data/processed", exist_ok=True)

with open(OUT_JSON, "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=2, default=str)

with open(OUT_MD, "w", encoding="utf-8") as f:
    f.write("# 1405.xlsx — Workbook Profile\n\n")
    f.write(
        f"- Generated: {out['generated_at_utc']}\n"
        f"- Sheets: **{out['sheet_count']}**\n\n"
    )

    for s in out["sheets"]:
        f.write(f"## {s['sheet_name']}\n\n")
        f.write(
            f"- Dimensions: {s['max_row']} × {s['max_column']}\n"
            f"- Non-empty rows: {s['nonempty_rows']}\n"
            f"- Formulas: {s['formula_count']}\n"
            f"- Merged ranges: {s['merged_range_count']}\n"
            f"- Header row: {s['provisional_header_row']}\n\n"
        )
        f.write(
            "| Col | Header | Nonblank | Blank | Types | Samples |\n"
            "|---:|---|---:|---:|---|---|\n"
        )

        for c in s["columns"]:
            samples_text = ", ".join(
                c["sample_unique_values"][:8]
            ).replace("|", "\\|")

            f.write(
                f"| {c['column_index']} | "
                f"{c['provisional_header']!r} | "
                f"{c['nonblank_count']} | "
                f"{c['blank_count']} | "
                f"{c['data_types']} | "
                f"{samples_text} |\n"
            )
