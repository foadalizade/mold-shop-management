#!/usr/bin/env python3
"""
Build AI-readable data from the current 1405.xlsx source.

Priority:
1) ./1405.xlsx  (the file currently maintained by the existing Auto Sync)
2) ./data/source/1405.xlsx (fallback)

Raw Excel is never modified.
This script only writes data/processed/ai/1405/.
RepairCase detection is intentionally NOT finalized here; this layer prepares
stable, text-readable inputs for the next analysis stage.
"""

from __future__ import annotations

import csv
import json
import math
import os
import re
from collections import defaultdict
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
PRIMARY_SOURCE = ROOT / "1405.xlsx"
FALLBACK_SOURCE = ROOT / "data" / "source" / "1405.xlsx"
OUTPUT_DIR = ROOT / "data" / "processed" / "ai" / "1405"

RAW_FIELDS = [
    ("F001", "ردیف", "RecordID"),
    ("F002", "قالب / قطعه / دستگاه", "AssetName"),
    ("F003", "نوع تعمیر", "WorkType"),
    ("F004", "اپراتور", "OperatorID"),
    ("F005", "مشتری", "CustomerID"),
    ("F006", "واحد درخواست کننده", "RequestingUnitID"),
    ("F007", "زیرگروه", "SubGroup"),
    ("F008", "کد قالب", "MoldCode"),
    ("F009", "شماره نامه درخواست", "RequestLetterNo"),
    ("F010", "خرابی", "FailureDescriptionRaw"),
    ("F011", "تعداد", "Quantity"),
    ("F012", "درخواست کننده", "RequesterID"),
    ("F013", "کد دستگاه خراب", "RelatedMachineCode"),
    ("F014", "ایستگاه", "StationID"),
    ("F015", "دستگاه", "MachineName"),
    ("F016", "قطعه", "PartName"),
    ("F017", "Part Code", "PartCode"),
    ("F018", "شرح عملیات", "OperationDescriptionRaw"),
    ("F019", "نتیجه فرایند", "ProcessResultRaw"),
    ("F020", "تاریخ ارسال به سختکاری", "HardeningSentDate"),
    ("F021", "مجری سختکاری", "HardeningExecutor"),
    ("F022", "هزینه سختکاری تومان", "HardeningCost"),
    ("F023", "تاریخ دریافت", "HardeningReceivedDate"),
    ("F024", "تاریخ ارسال به کروم", "ChromeSentDate"),
    ("F025", "مجری کروم", "ChromeExecutor"),
    ("F026", "هزینه کروم تومان", "ChromeCost"),
    ("F027", "تاریخ دریافت", "ChromeReceivedDate"),
    ("F028", "ماه", "Period"),
    ("F029", "تاریخ", "WorkDate"),
    ("F030", "ساعت کاری", "WorkShift"),
    ("F031", "مقدار دقیقه کار شده", "WorkedMinutes"),
    ("F032", "مقدار ساعت کار شده", "WorkedHours"),
    ("F033", "مرخصی اپراتور", "OperatorLeave"),
    ("F034", "ساعت مرخصی اپراتور", "OperatorLeaveHours"),
    ("F035", "توضیحات", "Notes"),
    ("F036", "قطعات خراب شده", "DamagedParts"),
    ("F037", "علت خرابی ثبت شده2", "RecordedFailureCause"),
    ("F038", "تاریخ3", "AdditionalProcessDateTime"),
    ("F039", "تایید کننده", "ApproverID"),
    ("F040", "تاریخ تحویل تعمیر/استارت تولید", "RepairDeliveryOrProductionStartDate"),
    ("F041", "تاریخ تحویل قالب به تولید", "MoldDeliveryToProductionDate"),
    ("F042", "فرم صدور تایید", "ApprovalFormStatus"),
    ("F043", "فی (تومان)", "UnitPrice"),
    ("F044", "قیمت کار شده (تومان)", "WorkPrice"),
    ("F045", "دقیقه خرابی دستگاه تراشکاری", "LatheMachineDowntimeMinutes"),
    ("F046", "ساعت خرابی دستگاه تراشکاری", "LatheMachineDowntimeHours"),
    ("F047", "", "UnnamedSourceField"),
]

STANDARD_HEADERS = [x[2] for x in RAW_FIELDS]
OUTPUT_HEADERS = (
    STANDARD_HEADERS
    + [
        "SourceSheet",
        "SourceRowNumber",
        "RecordKey",
        "AssetType",
        "MoldCodeNorm",
        "RequestLetterNoNorm",
        "WorkTypeNorm",
        "ProcessStatus",
    ]
)


def clean_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    text = str(value)
    text = text.translate(str.maketrans(
        "۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩",
        "01234567890123456789",
    ))
    text = text.replace("\u200c", " ").replace("\ufeff", "")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def norm_code(value) -> str:
    text = clean_text(value)
    text = text.replace("ي", "ی").replace("ك", "ک")
    return text.upper()


def asset_type(value) -> str:
    text = clean_text(value)
    if "قالب" in text:
        return "قالب"
    if "قطعه" in text:
        return "قطعه"
    if "دستگاه" in text:
        return "دستگاه"
    return "نامشخص"


def process_status(value) -> str:
    text = clean_text(value).lower()
    if text in {"ok", "ok."}:
        return "OK"
    if text in {"nok", "nc", "غیر قابل تعمیر"}:
        return "NOK"
    if not text:
        return ""
    return "UNKNOWN"


def find_header_row(ws) -> int | None:
    for row_number in range(1, min(ws.max_row, 20) + 1):
        values = {clean_text(v) for v in next(ws.iter_rows(
            min_row=row_number,
            max_row=row_number,
            values_only=True,
        ))}
        if "ردیف" in values and "نوع تعمیر" in values:
            return row_number
    return None


def build_record(sheet_name: str, row_number: int, values: list) -> dict | None:
    if len(values) < 47:
        values = values + [None] * (47 - len(values))
    values = values[:47]

    cleaned = [clean_text(v) for v in values]
    if not cleaned[0]:
        return None

    row = {standard: cleaned[i] for i, (_, _, standard) in enumerate(RAW_FIELDS)}
    row["SourceSheet"] = sheet_name
    row["SourceRowNumber"] = str(row_number)

    record_id = row["RecordID"]
    period = row["Period"] or sheet_name
    row["RecordKey"] = f"1405-{clean_text(period)}-{record_id}"

    row["AssetType"] = asset_type(row["AssetName"])
    row["MoldCodeNorm"] = norm_code(row["MoldCode"])
    row["RequestLetterNoNorm"] = norm_code(row["RequestLetterNo"])
    row["WorkTypeNorm"] = clean_text(row["WorkType"])
    row["ProcessStatus"] = process_status(row["ProcessResultRaw"])
    return row


def write_csv(path: Path, rows: list[dict], headers: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    source = PRIMARY_SOURCE if PRIMARY_SOURCE.exists() else FALLBACK_SOURCE
    if not source.exists():
        raise FileNotFoundError(
            f"1405.xlsx not found in either: {PRIMARY_SOURCE} or {FALLBACK_SOURCE}"
        )

    workbook = load_workbook(source, read_only=True, data_only=True)
    all_rows: list[dict] = []
    sheet_summary: list[dict] = []
    sheet_outputs: dict[str, list[dict]] = defaultdict(list)

    for sheet_name in workbook.sheetnames:
        if sheet_name == "Total":
            continue

        ws = workbook[sheet_name]
        header_row = find_header_row(ws)

        if header_row is None:
            sheet_summary.append({
                "SheetName": sheet_name,
                "HeaderRow": "",
                "OperationalRecords": 0,
                "Status": "HEADER_NOT_FOUND",
            })
            continue

        count = 0
        for row_number, values in enumerate(
            ws.iter_rows(
                min_row=header_row + 1,
                max_row=ws.max_row,
                max_col=47,
                values_only=True,
            ),
            start=header_row + 1,
        ):
            record = build_record(sheet_name, row_number, list(values))
            if record is None:
                continue
            all_rows.append(record)
            sheet_outputs[sheet_name].append(record)
            count += 1

        sheet_summary.append({
            "SheetName": sheet_name,
            "HeaderRow": header_row,
            "OperationalRecords": count,
            "Status": "OK",
        })

    workbook.close()

    write_csv(OUTPUT_DIR / "operational_records.csv", all_rows, OUTPUT_HEADERS)
    write_csv(
        OUTPUT_DIR / "sheet_summary.csv",
        sheet_summary,
        ["SheetName", "HeaderRow", "OperationalRecords", "Status"],
    )

    # Per-sheet AI-readable files.
    for sheet_name, rows in sheet_outputs.items():
        safe_name = sheet_name.replace("/", "_").replace("\\", "_")
        write_csv(OUTPUT_DIR / f"{safe_name}.csv", rows, OUTPUT_HEADERS)

    mold_rows = [
        row for row in all_rows
        if row["AssetType"] == "قالب" and row["MoldCodeNorm"]
    ]
    repair_candidates = [
        row for row in mold_rows
        if "تعمیر" in row["WorkTypeNorm"]
    ]

    write_csv(OUTPUT_DIR / "mold_records.csv", mold_rows, OUTPUT_HEADERS)
    write_csv(
        OUTPUT_DIR / "mold_repair_candidates.csv",
        repair_candidates,
        OUTPUT_HEADERS,
    )

    # Candidate-level summary, deliberately not a final RepairCase engine.
    grouped: dict[str, list[dict]] = defaultdict(list)
    for row in repair_candidates:
        grouped[row["MoldCodeNorm"]].append(row)

    mold_summary: list[dict] = []
    for mold_code, rows in sorted(grouped.items()):
        request_numbers = sorted({
            row["RequestLetterNoNorm"]
            for row in rows
            if row["RequestLetterNoNorm"]
        })
        work_dates = [
            row["WorkDate"] for row in rows if row["WorkDate"]
        ]
        mold_summary.append({
            "MoldCode": rows[0]["MoldCode"],
            "MoldCodeNorm": mold_code,
            "MoldName": rows[0]["AssetName"],
            "RepairLikeRecordCount": len(rows),
            "DistinctRequestLetterCount": len(request_numbers),
            "RequestLetterNos": " | ".join(request_numbers),
            "FirstWorkDate": min(work_dates) if work_dates else "",
            "LastWorkDate": max(work_dates) if work_dates else "",
        })

    write_csv(
        OUTPUT_DIR / "mold_summary.csv",
        mold_summary,
        [
            "MoldCode",
            "MoldCodeNorm",
            "MoldName",
            "RepairLikeRecordCount",
            "DistinctRequestLetterCount",
            "RequestLetterNos",
            "FirstWorkDate",
            "LastWorkDate",
        ],
    )

    metadata = {
        "source_file": str(source.relative_to(ROOT)).replace("\\", "/"),
        "github_commit_sha": os.environ.get("GITHUB_SHA", ""),
        "generated_at_utc": datetime.utcnow().replace(microsecond=0).isoformat() + "Z",
        "sheet_count_excluding_total": len(sheet_outputs),
        "operational_record_count": len(all_rows),
        "mold_record_count": len(mold_rows),
        "mold_repair_candidate_count": len(repair_candidates),
        "output_role": "AI-readable staging layer; RepairCase logic not final",
        "raw_data_preserved": True,
    }

    (OUTPUT_DIR / "index.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    schema = {
        "version": "ai-readable-1405-v1",
        "standard_fields": [
            {
                "source_field": source_id,
                "raw_name": raw_name,
                "standard_name": standard_name,
            }
            for source_id, raw_name, standard_name in RAW_FIELDS
        ],
        "derived_fields": [
            "SourceSheet",
            "SourceRowNumber",
            "RecordKey",
            "AssetType",
            "MoldCodeNorm",
            "RequestLetterNoNorm",
            "WorkTypeNorm",
            "ProcessStatus",
        ],
    }

    (OUTPUT_DIR / "schema.json").write_text(
        json.dumps(schema, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    print(f"Source: {source}")
    print(f"Operational records: {len(all_rows)}")
    print(f"Mold records: {len(mold_rows)}")
    print(f"Mold repair candidates: {len(repair_candidates)}")
    print(f"Output: {OUTPUT_DIR}")


if __name__ == "__main__":
    main()
